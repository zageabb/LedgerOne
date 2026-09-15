from __future__ import annotations

from collections import Counter
from io import BytesIO
import re
from pathlib import Path

from flask import current_app

from ledgerone.extensions import db
from ledgerone.models.core import utcnow
from ledgerone.modules.ai.models import KnowledgeChunk, KnowledgeSource
from ledgerone.services.audit import record_audit_event
from ledgerone.services.context import AccessContext


class KnowledgeError(RuntimeError):
    pass


_TOKEN_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9_\-]{1,}")
_STOPWORDS = {
    "about", "after", "again", "also", "and", "are", "been", "before", "being",
    "between", "but", "can", "could", "does", "for", "from", "have", "how",
    "into", "its", "more", "not", "our", "should", "that", "the", "their",
    "then", "there", "these", "they", "this", "those", "through", "was", "what",
    "when", "where", "which", "with", "would", "you", "your",
}


def _identity_id(context: AccessContext) -> str | None:
    return context.user_id or context.api_key_id


def _require(context: AccessContext, permission: str):
    if not context.can(permission):
        raise PermissionError(f"Missing permission: {permission}")


def _tokens(text: str) -> list[str]:
    return [
        token.lower()
        for token in _TOKEN_RE.findall(text or "")
        if len(token) > 2 and token.lower() not in _STOPWORDS
    ]


class KnowledgeService:
    CHUNK_WORDS = 180
    CHUNK_OVERLAP = 35

    @classmethod
    def list_sources(cls, context: AccessContext):
        if not (context.can("ai.knowledge.read") or context.can("ai.knowledge.manage")):
            raise PermissionError("Missing permission: ai.knowledge.read")
        return (
            KnowledgeSource.query.filter_by(organisation_id=context.organisation_id)
            .order_by(KnowledgeSource.updated_at.desc(), KnowledgeSource.title.asc())
            .all()
        )

    @classmethod
    def _chunks(cls, text: str) -> list[str]:
        words = (text or "").split()
        if not words:
            return []
        size = cls.CHUNK_WORDS
        step = max(1, size - cls.CHUNK_OVERLAP)
        return [
            " ".join(words[start:start + size]).strip()
            for start in range(0, len(words), step)
            if words[start:start + size]
        ]

    @classmethod
    def _replace_chunks(cls, source: KnowledgeSource) -> int:
        KnowledgeChunk.query.filter_by(
            organisation_id=source.organisation_id,
            source_id=source.id,
        ).delete(synchronize_session=False)
        chunks = cls._chunks(source.content)
        for index, content in enumerate(chunks):
            db.session.add(
                KnowledgeChunk(
                    organisation_id=source.organisation_id,
                    source_id=source.id,
                    chunk_index=index,
                    content=content,
                )
            )
        source.status = "ready" if chunks else "empty"
        source.updated_at = utcnow()
        return len(chunks)

    @classmethod
    def create_text_source(
        cls,
        context: AccessContext,
        *,
        title: str,
        content: str,
        filename: str | None = None,
        media_type: str | None = "text/plain",
    ) -> KnowledgeSource:
        _require(context, "ai.knowledge.manage")
        title = (title or "").strip()
        content = (content or "").strip()
        if not title:
            raise KnowledgeError("Knowledge title is required")
        if not content:
            raise KnowledgeError("Knowledge source is empty")
        max_bytes = int(current_app.config.get("KNOWLEDGE_MAX_BYTES", 10 * 1024 * 1024))
        if len(content.encode("utf-8")) > max_bytes:
            raise KnowledgeError(f"Knowledge source exceeds the {max_bytes // (1024 * 1024)} MB limit")

        source = KnowledgeSource(
            organisation_id=context.organisation_id,
            title=title[:255],
            filename=(filename or "")[:255] or None,
            media_type=(media_type or "")[:120] or None,
            content=content,
            enabled=True,
            status="indexing",
            uploaded_by_type=context.identity_type,
            uploaded_by_id=_identity_id(context),
        )
        db.session.add(source)
        db.session.flush()
        chunk_count = cls._replace_chunks(source)
        record_audit_event(
            context,
            module_id="ai",
            action="knowledge_create",
            entity_type="knowledge_source",
            entity_id=source.id,
            detail={"title": source.title, "filename": source.filename, "chunks": chunk_count},
        )
        db.session.commit()
        return source

    @classmethod
    def _extract_upload(cls, filename: str, data: bytes) -> tuple[str, str]:
        suffix = Path(filename or "").suffix.lower()
        if suffix in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log"}:
            return data.decode("utf-8", errors="replace"), "text/plain"

        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise KnowledgeError("PDF Knowledge support requires pypdf") from exc
            try:
                reader = PdfReader(BytesIO(data))
                text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
            except Exception as exc:
                raise KnowledgeError(f"Could not read PDF Knowledge source: {exc}") from exc
            return text, "application/pdf"

        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise KnowledgeError("Word Knowledge support requires python-docx") from exc
            try:
                document = Document(BytesIO(data))
                parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            except Exception as exc:
                raise KnowledgeError(f"Could not read Word Knowledge source: {exc}") from exc
            return "\n".join(parts), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        raise KnowledgeError("Unsupported Knowledge file type. Use TXT, MD, CSV, JSON, YAML, PDF or DOCX.")

    @classmethod
    def create_upload(cls, context: AccessContext, uploaded_file, *, title: str | None = None):
        _require(context, "ai.knowledge.manage")
        filename = (uploaded_file.filename or "").strip()
        if not filename:
            raise KnowledgeError("A file is required")
        max_bytes = int(current_app.config.get("KNOWLEDGE_MAX_BYTES", 10 * 1024 * 1024))
        data = uploaded_file.stream.read(max_bytes + 1)
        if len(data) > max_bytes:
            raise KnowledgeError(f"Knowledge file exceeds the {max_bytes // (1024 * 1024)} MB limit")
        text, media_type = cls._extract_upload(filename, data)
        return cls.create_text_source(
            context,
            title=(title or Path(filename).stem or filename),
            content=text,
            filename=filename,
            media_type=media_type,
        )

    @classmethod
    def reindex(cls, context: AccessContext, source_id: str):
        _require(context, "ai.knowledge.manage")
        source = cls._get_source(context, source_id)
        cls._replace_chunks(source)
        record_audit_event(
            context,
            module_id="ai",
            action="knowledge_reindex",
            entity_type="knowledge_source",
            entity_id=source.id,
            detail={"title": source.title},
        )
        db.session.commit()
        return source

    @classmethod
    def set_enabled(cls, context: AccessContext, source_id: str, enabled: bool):
        _require(context, "ai.knowledge.manage")
        source = cls._get_source(context, source_id)
        source.enabled = bool(enabled)
        source.updated_at = utcnow()
        record_audit_event(
            context,
            module_id="ai",
            action="knowledge_enable" if enabled else "knowledge_disable",
            entity_type="knowledge_source",
            entity_id=source.id,
            detail={"title": source.title},
        )
        db.session.commit()
        return source

    @classmethod
    def delete(cls, context: AccessContext, source_id: str):
        _require(context, "ai.knowledge.manage")
        source = cls._get_source(context, source_id)
        title = source.title
        db.session.delete(source)
        record_audit_event(
            context,
            module_id="ai",
            action="knowledge_delete",
            entity_type="knowledge_source",
            entity_id=source_id,
            detail={"title": title},
        )
        db.session.commit()

    @staticmethod
    def _get_source(context: AccessContext, source_id: str):
        source = KnowledgeSource.query.filter_by(
            id=source_id,
            organisation_id=context.organisation_id,
        ).first()
        if not source:
            raise KnowledgeError("Knowledge source not found")
        return source

    @classmethod
    def retrieve(cls, context: AccessContext, query: str, *, limit: int = 4) -> list[dict]:
        _require(context, "ai.knowledge.read")
        query_tokens = _tokens(query)
        if not query_tokens:
            return []

        query_counts = Counter(query_tokens)
        rows = (
            db.session.query(KnowledgeChunk, KnowledgeSource)
            .join(KnowledgeSource, KnowledgeSource.id == KnowledgeChunk.source_id)
            .filter(
                KnowledgeChunk.organisation_id == context.organisation_id,
                KnowledgeSource.organisation_id == context.organisation_id,
                KnowledgeSource.enabled.is_(True),
                KnowledgeSource.status == "ready",
            )
            .all()
        )

        ranked: list[tuple[float, KnowledgeChunk, KnowledgeSource]] = []
        for chunk, source in rows:
            content_counts = Counter(_tokens(chunk.content))
            title_counts = Counter(_tokens(source.title))
            overlap = sum(min(count, content_counts.get(token, 0)) for token, count in query_counts.items())
            title_overlap = sum(min(count, title_counts.get(token, 0)) for token, count in query_counts.items())
            if overlap <= 0 and title_overlap <= 0:
                continue
            score = float(overlap + (title_overlap * 3))
            ranked.append((score, chunk, source))

        ranked.sort(key=lambda row: (-row[0], row[2].title.lower(), row[1].chunk_index))
        return [
            {
                "source_id": source.id,
                "title": source.title,
                "filename": source.filename,
                "chunk_index": chunk.chunk_index,
                "score": score,
                "excerpt": chunk.content,
            }
            for score, chunk, source in ranked[: max(1, min(int(limit), 8))]
        ]
